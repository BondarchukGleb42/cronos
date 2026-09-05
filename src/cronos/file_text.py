"""Read beyond the stored preview without imposing a document-length product cap."""

import io
from pathlib import Path

from cronos.artifacts import _decode, _table_text


def text_window(artifact, offset, length):
    extracted = artifact["extracted"]
    text = extracted.get("text", "")
    if extracted.get("text_truncated"):
        path = Path(artifact["path"])
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".csv"}:
            text, _ = _decode(path.read_bytes())
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            from docx import Document

            document = Document(io.BytesIO(path.read_bytes()))
            pieces = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                pieces.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
            text = "\n".join(pieces)
        elif extracted.get("tables"):
            text = "\n\n".join(_table_text(table) for table in extracted["tables"])
    return {
        "text": text[offset : offset + length],
        "total_chars": len(text),
        "next_offset": offset + length if offset + length < len(text) else None,
    }
