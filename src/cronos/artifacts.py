"""Synchronous, user-scoped document storage and deterministic file operations.

Call these CPU/file operations in a worker. Uploaded documents are untrusted data:
spreadsheet formulas and document macros are never executed by this module.
"""

import csv
import errno
import hashlib
import io
import json
import mimetypes
import os
import re
import tempfile
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from threading import Lock
from uuid import uuid4
from xml.sax.saxutils import escape

# This is a context preview limit, not an upload, document, or storage limit.
TEXT_PREVIEW_CHARS = 250_000
_FONT_LOCK = Lock()
_MIMES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".md": "text/markdown",
}
_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}


def _safe_filename(filename: str) -> str:
    name = filename.replace("\\", "/").rsplit("/", 1)[-1]
    name = re.sub(r"[\x00-\x1f\x7f]", "_", name).strip()
    return name if name not in {"", ".", ".."} else "artifact"


def _json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    # Array formula and rich-text objects also need a stable JSON representation.
    return getattr(value, "text", str(value))


def _display(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(_json_value(value))


def _preview(text: str) -> dict:
    return {
        "text": text[:TEXT_PREVIEW_CHARS],
        "text_truncated": len(text) > TEXT_PREVIEW_CHARS,
        "text_total_chars": len(text),
    }


def _decode(data: bytes) -> tuple[str, str]:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "utf-16"
    try:
        return data.decode("utf-8-sig"), "utf-8"
    except UnicodeDecodeError:
        # Windows-1251 is a common encoding of Russian CSV/TXT exports.
        try:
            return data.decode("cp1251"), "windows-1251"
        except UnicodeDecodeError as exc:
            raise ValueError(
                "Не удалось определить кодировку файла; сохраните его в UTF-8"
            ) from exc


def _table_text(table: dict) -> str:
    rows = [table["columns"], *table["rows"]]
    return "\n".join("\t".join(_display(value) for value in row) for row in rows)


def _extract(data: bytes, suffix: str) -> dict:
    if suffix in {".txt", ".md"}:
        text, encoding = _decode(data)
        return {**_preview(text), "encoding": encoding}
    if suffix == ".csv":
        text, encoding = _decode(data)
        try:
            dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        records = list(csv.reader(io.StringIO(text, newline=""), dialect=dialect))
        table = {"name": "CSV", "columns": records[0] if records else [], "rows": records[1:]}
        return {**_preview(text), "encoding": encoding, "tables": [table]}
    if suffix == ".xlsx":
        return _extract_xlsx(data)
    if suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(io.BytesIO(data))
        paragraphs, tables = [], []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                paragraphs.append(block.text)
            elif isinstance(block, Table):
                records = [[cell.text for cell in row.cells] for row in block.rows]
                table = {
                    "name": f"Table {len(tables) + 1}",
                    "columns": records[0] if records else [],
                    "rows": records[1:],
                }
                tables.append(table)
                paragraphs.append(_table_text(table))
        return {**_preview("\n".join(paragraphs)), "tables": tables}
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF защищен паролем; отправьте разблокированную копию")
        pages, texts, remaining = [], [], TEXT_PREVIEW_CHARS
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            texts.append(text)
            pages.append(
                {
                    "page": index,
                    "text": text[:remaining],
                    "text_truncated": len(text) > remaining,
                    "text_total_chars": len(text),
                    "needs_ocr": not bool(text.strip()),
                    "width": float(page.mediabox.width),
                    "height": float(page.mediabox.height),
                }
            )
            remaining = max(0, remaining - len(text))
        return {
            **_preview("\n\n".join(texts)),
            "pages": pages,
            "page_count": len(pages),
            "needs_ocr": any(page["needs_ocr"] for page in pages),
            "ocr_pages": [page["page"] for page in pages if page["needs_ocr"]],
        }
    if suffix in _IMAGE_EXTENSIONS:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as picture:
            metadata = {
                "width": picture.width,
                "height": picture.height,
                "format": picture.format,
                "mode": picture.mode,
                "frames": getattr(picture, "n_frames", 1),
            }
            picture.verify()
        return {"text": "", "image": metadata, "needs_vision": True}
    return {"text": "", "supported": False, "reason": "Нет обработчика этого формата"}


def _extract_xlsx(data: bytes) -> dict:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    tables, formulas = [], []
    try:
        for sheet in book:
            # Spreadsheet dimensions may be wrong in exports from third-party software.
            sheet.reset_dimensions()
            cached_sheet = cached[sheet.title]
            cached_sheet.reset_dimensions()
            cached_rows = iter(cached_sheet.iter_rows())
            records = []
            for row in sheet.iter_rows():
                cached_row = next(cached_rows, ())
                values = []
                for index, cell in enumerate(row):
                    values.append(_json_value(cell.value))
                    if cell.data_type == "f":
                        cached_value = cached_row[index].value if index < len(cached_row) else None
                        formulas.append(
                            {
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                                "formula": _json_value(cell.value),
                                "cached_value": _json_value(cached_value),
                            }
                        )
                records.append(values)
            tables.append(
                {
                    "name": sheet.title,
                    "columns": records[0] if records else [],
                    "rows": records[1:],
                }
            )
    finally:
        book.close()
        cached.close()
    return {
        **_preview("\n\n".join(table["name"] + "\n" + _table_text(table) for table in tables)),
        "tables": tables,
        "formulas": formulas,
        "formulas_evaluated": False,
    }


def _csv_value(value):
    if not isinstance(value, str):
        return _display(value)
    # Quoting CSV fields does not prevent Excel from evaluating them on opening.
    if value.lstrip(" \t\r\n\ufeff").startswith(("=", "+", "-", "@")) or value.startswith(
        ("\t", "\r", "\n")
    ):
        return "'" + value
    return value


def _font_name() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    with _FONT_LOCK:
        name = "CronosSans"
        if name not in pdfmetrics.getRegisteredFontNames():
            candidates = [
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                "/usr/share/fonts/dejavu/DejaVuSans.ttf",
                "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
                "/System/Library/Fonts/Supplemental/Arial.ttf",
                "/Library/Fonts/Arial.ttf",
            ]
            path = next((item for item in candidates if Path(item).is_file()), None)
            if path is None:
                raise RuntimeError("Для PDF нужен шрифт с кириллицей: установите fonts-dejavu-core")
            pdfmetrics.registerFont(TTFont(name, path))
        return name


def _generate_pdf(content: str, columns: list[str], rows: list[list]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    target = io.BytesIO()
    count = max([len(columns), *(len(row) for row in rows)], default=0)
    page_size = landscape(A4) if count > 5 else A4
    document = SimpleDocTemplate(
        target,
        pagesize=page_size,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    font = _font_name()
    style = ParagraphStyle("Body", fontName=font, fontSize=10, leading=14, spaceAfter=7)
    table_style = ParagraphStyle(
        "Cell",
        fontName=font,
        fontSize=9,
        leading=12,
        alignment=TA_LEFT,
        splitLongWords=True,
    )
    story = []
    for line in content.splitlines():
        story.append(Paragraph(escape(line), style) if line else Spacer(1, 7))
    if count:
        if story:
            story.append(Spacer(1, 8))
        records = ([columns] if columns else []) + rows
        cells = [
            [
                Paragraph(escape(_display(value)).replace("\n", "<br/>"), table_style)
                for value in [*row, *([None] * (count - len(row)))]
            ]
            for row in records
        ]
        table = LongTable(
            cells,
            colWidths=[document.width / count] * count,
            repeatRows=1 if columns else 0,
            splitByRow=1,
            splitInRow=1,
        )
        commands = [
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D9D9")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]
        if columns:
            commands.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF2")))
        table.setStyle(TableStyle(commands))
        story.append(table)
    document.build(story or [Spacer(1, 1)])
    return target.getvalue()


def _generate_docx(content: str, columns: list[str], rows: list[list]) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.7)
    for line in content.splitlines():
        document.add_paragraph(line)
    count = max([len(columns), *(len(row) for row in rows)], default=0)
    if count:
        table = document.add_table(rows=0, cols=count)
        table.style = "Table Grid"
        for index, record in enumerate(([columns] if columns else []) + rows):
            cells = table.add_row().cells
            for cell, value in zip(cells, record, strict=False):
                cell.text = _display(value)
            if index == 0 and columns:
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for cell in cells:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E8EDF2")
                    cell._tc.get_or_add_tcPr().append(shade)
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
    target = io.BytesIO()
    document.save(target)
    return target.getvalue()


def _generate_xlsx(content: str, columns: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "Данные"
    records = ([columns] if columns else []) + rows
    if not records:
        records = [
            [part]
            for line in content.splitlines()
            for part in [line[index : index + 32767] for index in range(0, len(line), 32767)]
            or [""]
        ]
        records = records or [[content]]
    for row in records:
        values = [
            json.dumps(item, ensure_ascii=False) if isinstance(item, (list, dict)) else item
            for item in row
        ]
        if any(isinstance(value, str) and len(value) > 32767 for value in values):
            raise ValueError(
                "Ячейка XLSX вмещает не более 32767 символов; используйте CSV для этих данных"
            )
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                # Untrusted formulas stay actual strings, without changing their text.
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if columns:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EDF2")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    for index, cells in enumerate(sheet.columns, start=1):
        length = max((len(_display(cell.value)) for cell in cells), default=10)
        sheet.column_dimensions[get_column_letter(index)].width = min(60, max(12, length + 2))
    if content and (columns or rows):
        notes = book.create_sheet("Описание")
        for line in content.splitlines():
            for part in [line[index : index + 32767] for index in range(0, len(line), 32767)] or [
                ""
            ]:
                notes.append([part])
                notes.cell(notes.max_row, 1).data_type = "s"
                notes.cell(notes.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        notes.column_dimensions["A"].width = 80
    target = io.BytesIO()
    book.save(target)
    book.close()
    return target.getvalue()


def render_pdf_pages(path: str, max_pages: int | None = None) -> dict:
    """Render validated PDF paths to PNG bytes, explicitly indicating partial output.

    The caller must authorize the path with ArtifactManager.read first. No page
    limit is applied unless the caller supplies max_pages for its own workflow.
    """
    import pypdfium2 as pdfium

    if max_pages is not None and (
        isinstance(max_pages, bool) or not isinstance(max_pages, int) or max_pages < 1
    ):
        raise ValueError("max_pages должен быть положительным числом или None")
    images = []
    with pdfium.PdfDocument(path) as document:
        total = len(document)
        count = min(total, max_pages) if max_pages is not None else total
        for index in range(count):
            page = document[index]
            try:
                bitmap = page.render(scale=2)
                try:
                    with bitmap.to_pil() as picture:
                        output = io.BytesIO()
                        picture.save(output, format="PNG")
                        images.append(output.getvalue())
                finally:
                    bitmap.close()
            finally:
                page.close()
    return {
        "images": images,
        "page_numbers": list(range(1, count + 1)),
        "total_pages": total,
        "rendered_pages": count,
        "partial": count < total,
    }


class ArtifactManager:
    def __init__(self, root: str):
        self.root = Path(root).expanduser().resolve()
        self.root.mkdir(parents=True, exist_ok=True, mode=0o700)

    def _user_dir(self, user_id: int) -> Path:
        if isinstance(user_id, bool) or not isinstance(user_id, int) or user_id <= 0:
            raise ValueError("user_id должен быть положительным целым числом")
        directory = self.root / str(user_id)
        if directory.is_symlink():
            raise PermissionError("Символические ссылки в каталоге пользователя запрещены")
        directory.mkdir(mode=0o700, exist_ok=True)
        return directory

    def ingest(self, user_id: int, filename: str, data: bytes) -> dict:
        directory = self._user_dir(user_id)
        filename = _safe_filename(filename)
        suffix = Path(filename).suffix.lower()
        extracted = _extract(data, suffix)
        identifier = str(uuid4())
        # Extension is useful for native apps; original names never become paths.
        extension = suffix if re.fullmatch(r"\.[a-z0-9]{1,12}", suffix) else ""
        path = directory / (identifier + extension)
        fd, temporary = tempfile.mkstemp(prefix=".upload-", dir=directory)
        try:
            with os.fdopen(fd, "wb") as stream:
                stream.write(data)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, path)
            directory_fd = os.open(directory, os.O_RDONLY)
            try:
                try:
                    os.fsync(directory_fd)
                except OSError as exc:
                    if exc.errno not in {errno.EINVAL, errno.ENOTSUP}:
                        raise
            finally:
                os.close(directory_fd)
        finally:
            Path(temporary).unlink(missing_ok=True)
        return {
            "id": identifier,
            "path": str(path),
            "filename": filename,
            "mime": _MIMES.get(suffix)
            or mimetypes.guess_type(filename)[0]
            or "application/octet-stream",
            "size_bytes": len(data),
            "checksum": hashlib.sha256(data).hexdigest(),
            "extracted": extracted,
        }

    def read(self, user_id: int, path: str) -> dict:
        directory = self._user_dir(user_id)
        supplied = Path(path)
        if ".." in supplied.parts:
            raise PermissionError("Путь за пределы каталога пользователя запрещен")
        candidate = supplied if supplied.is_absolute() else directory / supplied
        if candidate.is_symlink() or candidate.resolve().parent != directory:
            raise PermissionError("Файл не принадлежит этому пользователю")
        if not candidate.is_file():
            raise FileNotFoundError("Файл не найден")
        return _extract(candidate.read_bytes(), candidate.suffix.lower())

    def generate(
        self,
        user_id: int,
        format: str,
        filename: str,
        content: str,
        columns: list[str] | None = None,
        rows: list[list] | None = None,
    ) -> dict:
        self._user_dir(user_id)
        format = format.lower().lstrip(".")
        columns, rows = columns or [], rows or []
        if rows and not columns:
            columns = [f"Столбец {index + 1}" for index in range(max(map(len, rows)))]
        if format == "pdf":
            data = _generate_pdf(content, columns, rows)
        elif format == "docx":
            data = _generate_docx(content, columns, rows)
        elif format == "xlsx":
            data = _generate_xlsx(content, columns, rows)
        elif format == "csv":
            stream = io.StringIO(newline="")
            writer = csv.writer(stream)
            records = ([columns] if columns else []) + rows
            if not records:
                records = [[line] for line in content.splitlines()] or [[content]]
            writer.writerows([_csv_value(value) for value in row] for row in records)
            data = stream.getvalue().encode("utf-8-sig")
        elif format == "txt":
            text = content
            if columns or rows:
                text += ("\n\n" if text else "") + _table_text({"columns": columns, "rows": rows})
            data = text.encode("utf-8")
        else:
            raise ValueError(f"Генерация формата {format!r} не поддерживается")
        name = str(Path(_safe_filename(filename)).with_suffix("." + format))
        return self.ingest(user_id, name, data)
