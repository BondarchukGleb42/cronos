import csv
import hashlib
import io
import json
import os
from datetime import date
from pathlib import Path
from uuid import UUID
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.errors import PdfReadError

from cronos.artifacts import ArtifactManager, render_pdf_pages


@pytest.fixture
def manager(tmp_path):
    return ArtifactManager(str(tmp_path / "artifacts"))


def test_ingest_is_durable_unique_and_user_scoped(manager):
    data = "Привет, мир!".encode()
    first = manager.ingest(42, "../../личное.txt", data)
    second = manager.ingest(42, "личное.txt", b"Second version")
    assert first["filename"] == "личное.txt"
    assert first["checksum"] == hashlib.sha256(data).hexdigest()
    assert first["size_bytes"] == len(data)
    assert first["mime"] == "text/plain"
    assert UUID(first["id"])
    assert first["path"] != second["path"]
    assert Path(first["path"]).parent == manager.root / "42"
    assert Path(first["path"]).read_bytes() == data
    assert os.stat(first["path"]).st_mode & 0o777 == 0o600
    restarted = ArtifactManager(str(manager.root))
    assert restarted.read(42, first["path"])["text"] == "Привет, мир!"
    assert not list((manager.root / "42").glob(".upload-*"))
    with pytest.raises(PermissionError):
        manager.read(43, first["path"])
    with pytest.raises(PermissionError):
        manager.read(43, "../42/" + Path(first["path"]).name)


def test_symlinks_cannot_escape_user_storage(manager, tmp_path):
    artifact = manager.ingest(42, "private.txt", b"private")
    manager._user_dir(43)
    link = manager.root / "43" / "leak.txt"
    link.symlink_to(artifact["path"])
    with pytest.raises(PermissionError):
        manager.read(43, str(link))
    (manager.root / "44").symlink_to(tmp_path, target_is_directory=True)
    with pytest.raises(PermissionError):
        manager.ingest(44, "escape.txt", b"data")


@pytest.mark.parametrize("user_id", [True, 0, -1, "../42"])
def test_invalid_user_identity_is_rejected(manager, user_id):
    with pytest.raises(ValueError):
        manager.ingest(user_id, "file.txt", b"value")


def test_explicit_preview_truncation_does_not_truncate_stored_file(manager, monkeypatch):
    monkeypatch.setattr("cronos.artifacts.TEXT_PREVIEW_CHARS", 8)
    artifact = manager.ingest(42, "long.txt", "Длинный русский текст".encode())
    assert artifact["extracted"]["text"] == "Длинный "
    assert artifact["extracted"]["text_truncated"] is True
    assert artifact["extracted"]["text_total_chars"] == len("Длинный русский текст")
    assert Path(artifact["path"]).read_text() == "Длинный русский текст"


def test_failed_atomic_write_leaves_no_partial_file(manager, monkeypatch):
    def fail_replace(*args):
        raise OSError("disk unavailable")

    monkeypatch.setattr("cronos.artifacts.os.replace", fail_replace)
    with pytest.raises(OSError, match="disk unavailable"):
        manager.ingest(42, "draft.txt", b"incomplete")
    assert list((manager.root / "42").iterdir()) == []


def test_csv_roundtrip_handles_unicode_numbers_quotes_and_formula_injection(manager):
    artifact = manager.generate(
        42,
        "csv",
        "таблица.csv",
        "",
        ["Описание", "Сумма"],
        [["Строка, с запятой\nи переносом", 12.5], ["=1+1", -4], [" \t@SUM(1)", 0]],
    )
    with Path(artifact["path"]).open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert rows[1] == ["Строка, с запятой\nи переносом", "12.5"]
    assert rows[2] == ["'=1+1", "-4"]
    assert rows[3][0] == "' \t@SUM(1)"
    assert artifact["extracted"]["tables"][0]["rows"] == rows[1:]


def test_windows_csv_is_detected_and_utf16_text_is_read(manager):
    artifact = manager.ingest(
        42, "экспорт.csv", "Товар;Цена\nМолоко;90\nХлеб;45\n".encode("cp1251")
    )
    assert artifact["extracted"]["encoding"] == "windows-1251"
    assert artifact["extracted"]["tables"][0]["columns"] == ["Товар", "Цена"]
    text = manager.ingest(42, "текст.txt", "Русский текст".encode("utf-16"))
    assert text["extracted"]["text"] == "Русский текст"


def test_xlsx_keeps_numeric_types_totals_dates_and_literal_formulas(manager):
    artifact = manager.generate(
        42,
        "xlsx",
        "закупки.bad",
        "Расчет закупок на неделю",
        ["Товар", "Количество", "Цена", "Сумма"],
        [
            ["Молоко", 2, 90, 180],
            ["Хлеб", 3, 45, 135],
            ["Итого", None, None, 315],
            ['=HYPERLINK("https://example.com")', date(2026, 9, 5), True, 0],
        ],
    )
    assert artifact["filename"] == "закупки.xlsx"
    book = load_workbook(artifact["path"])
    sheet = book["Данные"]
    assert sheet["B2"].value == 2 and sheet["B2"].data_type == "n"
    assert (
        sheet["D4"].value
        == sheet["B2"].value * sheet["C2"].value + sheet["B3"].value * sheet["C3"].value
    )
    assert sheet["A5"].data_type == "s"
    assert sheet["A5"].value.startswith("=HYPERLINK")
    assert sheet["B5"].is_date
    assert sheet.freeze_panes == "A2"
    assert book["Описание"]["A1"].value == "Расчет закупок на неделю"
    book.close()
    assert artifact["extracted"]["tables"][0]["rows"][0] == ["Молоко", 2, 90, 180]
    assert artifact["extracted"]["formulas"] == []
    json.dumps(artifact, ensure_ascii=False)


def test_imported_xlsx_formulas_are_reported_with_cache_and_never_executed(manager):
    book = Workbook()
    book.active.append(["Количество", "Цена", "Сумма"])
    book.active.append([3, 15, "=A2*B2"])
    target = io.BytesIO()
    book.save(target)
    # Supply a known Excel-calculated cached result to check both representations.
    patched = io.BytesIO()
    with ZipFile(io.BytesIO(target.getvalue())) as source, ZipFile(patched, "w") as output:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>A2*B2</f><v></v>", b"<f>A2*B2</f><v>45</v>")
            output.writestr(item, data)
    artifact = manager.ingest(42, "формулы.xlsx", patched.getvalue())
    assert artifact["extracted"]["tables"][0]["rows"] == [[3, 15, "=A2*B2"]]
    assert artifact["extracted"]["formulas"] == [
        {"sheet": "Sheet", "cell": "C2", "formula": "=A2*B2", "cached_value": 45},
    ]
    assert artifact["extracted"]["formulas_evaluated"] is False


def test_generation_preserves_first_data_row_when_columns_are_not_supplied(manager):
    artifact = manager.generate(42, "xlsx", "data", "", rows=[[10, 20], [30, 40]])
    assert artifact["extracted"]["tables"][0]["columns"] == ["Столбец 1", "Столбец 2"]
    assert artifact["extracted"]["tables"][0]["rows"] == [[10, 20], [30, 40]]


def test_xlsx_format_limit_never_silently_truncates_a_cell(manager):
    long_text = "а" * 40_000
    with pytest.raises(ValueError, match="32767"):
        manager.generate(42, "xlsx", "data", "", ["Текст"], [[long_text]])
    artifact = manager.generate(42, "xlsx", "text", long_text)
    table = artifact["extracted"]["tables"][0]
    assert table["columns"][0] + table["rows"][0][0] == long_text


@pytest.mark.parametrize("format", ["pdf", "docx", "txt"])
def test_document_roundtrip_cyrillic_table_and_reingest(manager, format):
    artifact = manager.generate(
        42,
        format,
        "Отчет",
        "Закупки на неделю\nПлан включает молоко, хлеб и фрукты.\nБюджет < 500 & > 300.",
        ["Продукт", "Количество", "Стоимость"],
        [["Молоко", 2, 180], ["Хлеб", 3, 135], ["Итого", "", 315]],
    )
    assert artifact["filename"] == f"Отчет.{format}"
    extracted = manager.read(42, artifact["path"])
    assert "Закупки на неделю" in extracted["text"]
    assert "Молоко" in extracted["text"]
    assert "315" in extracted["text"]
    assert "Бюджет < 500 & > 300." in extracted["text"]
    reuploaded = manager.ingest(43, artifact["filename"], Path(artifact["path"]).read_bytes())
    assert reuploaded["extracted"]["text"] == extracted["text"]
    if format == "pdf":
        assert extracted["needs_ocr"] is False
        assert len(PdfReader(artifact["path"]).pages) == 1


def test_image_metadata_and_scanned_pdf_ocr_detection(manager):
    picture = Image.new("RGB", (64, 48), "white")
    png = io.BytesIO()
    picture.save(png, format="PNG")
    image = manager.ingest(42, "фото.png", png.getvalue())
    assert image["extracted"]["image"]["width"] == 64
    assert image["extracted"]["image"]["height"] == 48
    assert image["extracted"]["needs_vision"] is True
    pdf = io.BytesIO()
    picture.save(pdf, format="PDF")
    scan = manager.ingest(42, "scan.pdf", pdf.getvalue())
    assert scan["extracted"]["needs_ocr"] is True
    assert scan["extracted"]["ocr_pages"] == [1]


def test_pdf_renderer_returns_all_pages_or_explicit_partial_result(manager):
    writer = PdfWriter()
    writer.add_blank_page(100, 200)
    writer.add_blank_page(100, 200)
    target = io.BytesIO()
    writer.write(target)
    artifact = manager.ingest(42, "two-pages.pdf", target.getvalue())
    all_pages = render_pdf_pages(artifact["path"])
    assert all_pages["total_pages"] == all_pages["rendered_pages"] == 2
    assert all_pages["partial"] is False
    assert len(all_pages["images"]) == 2
    assert all_pages["images"][0].startswith(b"\x89PNG\r\n\x1a\n")
    limited = render_pdf_pages(artifact["path"], max_pages=1)
    assert limited["partial"] is True
    assert limited["page_numbers"] == [1]
    assert limited["total_pages"] == 2


def test_invalid_document_fails_before_persisting_and_unknown_file_is_retained(manager):
    with pytest.raises(PdfReadError):
        manager.ingest(42, "broken.pdf", b"not a PDF")
    assert not list((manager.root / "42").iterdir())
    artifact = manager.ingest(42, "raw.bin", b"\x00\x01\x02")
    assert artifact["extracted"]["supported"] is False
    assert Path(artifact["path"]).read_bytes() == b"\x00\x01\x02"


def test_password_protected_pdf_is_explicitly_rejected(manager):
    writer = PdfWriter()
    writer.add_blank_page(100, 200)
    writer.encrypt("password")
    target = io.BytesIO()
    writer.write(target)
    with pytest.raises(ValueError, match="паролем"):
        manager.ingest(42, "locked.pdf", target.getvalue())
